# -*- coding: utf-8 -*-
"""uchu_build.py · BIM 4D de la SE Uchucchacua (habilitación del nuevo tablero CPM L-1123, servicio de AXON).

Arma `uchu4d.js` (bloque U4D) desde:
  · los controles diarios de AXON «Reporte_Diario__AXON_CPM_L-1123_<DD-Mes-AA>.xlsx» (uno por día; el más reciente trae
    la serie completa: hoja AVANCE, filas «Prog.» = % acumulado de la línea base de AXON —que sigue el cronograma V8 de
    CONENHUA— y «Ejec.» = % acumulado real al cierre de cada día; filas 54/55 = avance ponderado del servicio);
  · cada control del día para lo que es de ese día: cabecera (reporte N.º, día del plazo, ejecutado del día, personal,
    restricciones abiertas, habilitantes vencidos), causa y justificación por actividad, habilitantes y restricciones;
  · los cronogramas de referencia en PDF: «AGP-CNHA-138-UCHUC-CR V8» (CONENHUA) y «AX-26138 Cronograma Resumido» (AXON Rev. 1);
  · el índice de fotos ya difuminadas (`uchu_fotos.py`).
Solo cantidades, cargos y textos de obra: ningún nombre ni DNI del personal (los textos libres pasan por `sin_nombres`).

  python3 gen/uchu_build.py --reportes <carpeta con axon_AAAA-MM-DD.xlsx> --v8 <pdf> --axr1 <pdf> [--fotos fotos4d/fotos.json] --out uchu4d.js
"""
import argparse, datetime as dt, glob, json, os, re, subprocess, warnings
import openpyxl
warnings.filterwarnings('ignore')

def iso(v):
    if isinstance(v, dt.datetime): return v.date().isoformat()
    if isinstance(v, dt.date): return v.isoformat()
    return None

def num(v, d=6):   # 6 decimales: con 4, 0,545454 se mostraba 54,6 % en vez del 54,5 % del control
    return round(float(v), d) if isinstance(v, (int, float)) and not isinstance(v, bool) else None

NOMBRE = re.compile(r'(?:\b(?:el|la|al|del)\s+)?\b(Ing\.?|Ingeniero|Ingeniera|Sr\.?|Sra\.?|Srta\.?|Lic\.?|Téc\.?)\s+(?:[A-ZÁÉÍÓÚÑ](?:[a-záéíóúñ]+|\.)[ \t]*){1,4}')

def sin_nombres(t):
    """Quita nombres propios precedidos de un título («Ing. Iván Mejía» → «el ingeniero a cargo»)."""
    if t is None: return None
    t = str(t).strip()
    def cambio(m):
        art = re.match(r'(al|del)\s', m.group(0), re.I)      # «al Ing. X» → «al ingeniero a cargo»
        pre = (art.group(1) + ' ') if art else ('el ' if m.group(1).lower().startswith('ing') else 'la ')
        if art and art.group(1).lower() in ('al', 'del'): pre = art.group(1) + ' '
        return pre + ('ingeniero a cargo ' if m.group(1).lower().startswith('ing') else 'persona a cargo ')
    t = re.sub(r'\s{2,}', ' ', NOMBRE.sub(cambio, t)).strip()
    t = re.sub(r'\s+([.,;:])', r'\1', t)
    return t[:1].upper() + t[1:] if t else t

# ---------------------------------------------------------------- fotos ↔ actividades
# La hoja REGISTRO FOTOGRÁFICO nombra la actividad de algunas fotos (rel = «registro»). Para el resto, el pie de la foto
# se compara con estas palabras clave y la actividad solo se acepta si ese día se trabajó (o estaba programada): rel = «pie».
import unicodedata as _ud
def _n(t): return ''.join(c for c in _ud.normalize('NFD', str(t or '').lower()) if _ud.category(c) != 'Mn')
CLAVES = [
    (r'(marquillado|conexionado|timbrado).*ss\.? ?aa|ss\.? ?aa.*(marquillado|conexionado|timbrado)', ['12']),   # tableros de SS.AA.
    # 25-Oct-26: AXON amplió la 12 a la intervención de 3 tableros de SS.AA. (48 Vcc, 230 Vca esencial y 110 Vcc)
    (r'(breaker|interruptor|adecuacion|calado).*(ss\.? ?aa|servicios auxiliares)', ['12']),
    # 26-Oct-26: encendido y pruebas previas del tablero = 13; conexionado en el CPM = 11 (o 10 si es antes del 22-Oct)
    (r'encendido', ['13']),
    # 27-Oct-26: pruebas FAT/SAT de relés = 14; validación rojo-verde (continuidad) = timbrado del CPM o de SS.AA.
    (r'pruebas? (fat|sat)|(fat|sat)/(fat|sat)', ['14']),
    (r'rojo.?verde|validacion de (las )?conexiones', ['11', '12']),
    # 28-Oct-26: vistas del tablero «luego del acabado» = acabado de circuitos (17); entrega de bobinas sobrantes, carretes vacíos
    # y punta final de cada cable = orden, limpieza y desmovilización (18). Van antes de «bobina|carrete» (recepción y megado).
    (r'acabado|vista (frontal|posterior|interna|lateral)', ['17']),
    (r'sobrante|carrete vacio|punta final del cable', ['18']),
    (r'conexionado.*(e01|wr01|cpm)', ['11', '10']),
    (r'marquillado|timbrado', ['11']),                                                                         # tablero CPM
    (r'induccion|anexo 0?5|capacitacion', ['1']),
    (r'prueba.*aislamiento|aislamiento del cable|pulso de 1000|resistencia de los cables|megado', ['3']),
    (r'bobina|carrete', ['2', '3']),
    (r'tapas', ['5']),
    (r'sellad[oa].*(cable|multiconductor|punta|vulcaniz)|resguardo de (los )?cables|cubierta de plastico', ['6', '4']),   # sellado de puntas tras el tendido
    (r'tendido.*(tsm|ss\.? ?aa)', ['10.1']),
    (r'tendido.*(patio|tcmp|multiconductor)|movilizacion de cables|ordenamiento de cables en patio', ['6']),
    (r'desembalaje|anclaje|ingreso del tablero', ['7', '8']),
    (r'montaje de los equipos|equipos del tablero|\brele\b', ['9']),
    (r'registrador|tlrf', ['19']),
    (r'preparacion de cables? (para conexionado|(en|para) (el )?tablero)|conexionado del tablero', ['10']),
    (r'preparacion de cables|preparado de cable|vulcanizado|proteccion de la punta|codificacion', ['4']),
]
SSOMA = r'charla|orden y limpieza|segregacion|residuos|escalera|delimitacion'
def relacionar_fotos(fotos, acts, dias):
    por_n = {a['n']: a for a in acts}
    for f in fotos:
        if f.get('act'): f['rel'] = 'registro'; continue
        t = _n(f.get('pie')) + ' ' + _n(f.get('detalle'))
        if re.search(SSOMA, t): f['tema'] = 'SSOMA'; continue
        for pat, cands in CLAVES:
            if not re.search(pat, t): continue
            hechas = set((dias.get(f['f']) or {}).get('hechas') or [])
            ok = [n for n in cands if n in hechas] or [n for n in cands if por_n.get(n) and por_n[n]['ini'] and por_n[n]['ini'] <= f['f'] <= por_n[n]['fin']]
            if ok: f['act'] = ok[0]; f['rel'] = 'pie'
            break
    return fotos

# ---------------------------------------------------------------- control diario de AXON
def leer_avance(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    A = wb['AVANCE']
    fecha = iso(A.cell(1, 31).value)
    cab = {A.cell(3, c).value: A.cell(4, c).value for c in range(3, 40) if isinstance(A.cell(3, c).value, str)}
    cols = [(c, iso(A.cell(6, c).value)) for c in range(9, 32) if iso(A.cell(6, c).value)]
    acts, seccion, grupo = [], None, None
    for r in range(8, 54):
        if A.cell(r, 2).value: seccion = ' '.join(str(A.cell(r, 2).value).split())
        if A.cell(r, 3).value: grupo = ' '.join(str(A.cell(r, 3).value).split())
        if A.cell(r, 8).value != 'Prog.' or A.cell(r, 4).value in (None, ''): continue
        n = A.cell(r, 4).value
        n = str(int(n)) if isinstance(n, float) and n == int(n) else str(n)
        prog = {f: num(A.cell(r, c).value) for c, f in cols}
        ejec = {f: num(A.cell(r + 1, c).value) for c, f in cols} if A.cell(r + 1, 8).value == 'Ejec.' else {}
        # ritmo diario de la línea base = columna BW del control: MAX(I:AE de Prog.) / COUNT(I:AE de Prog.); con él el formato
        # vuelve a programar lo atrasado al día siguiente (columna CJ). BT = días de peso (G = BT ÷ Σ BT).
        vp = [v for v in prog.values() if v is not None]
        ritmo = round(max(vp) / len(vp), 6) if vp else 0
        acts.append(dict(n=n, nombre=' '.join(str(A.cell(r, 5).value).split()), ubic=str(A.cell(r, 6).value or '').strip(),
                         peso=num(A.cell(r, 7).value), seccion=seccion, grupo=grupo,
                         alcance='B' if (grupo or '').upper().startswith('ALCANCE B') else 'A',
                         pv=num(A.cell(r, 32).value), ev=num(A.cell(r, 33).value), estado=A.cell(r, 35).value,
                         causa=sin_nombres(A.cell(r, 37).value), justif=sin_nombres(A.cell(r, 38).value),
                         ini=iso(A.cell(r, 65).value), fin=iso(A.cell(r, 66).value), ritmo=ritmo, dias_peso=num(A.cell(r, 72).value),
                         _prog=prog, _ejec=ejec))
    tot = {}
    for r in range(50, 60):
        lab = str(A.cell(r, 2).value or '')
        if lab.startswith('AVANCE PONDERADO · PROGRAMADO'): tot['prog'] = {f: num(A.cell(r, c).value) for c, f in cols}
        if lab.startswith('AVANCE PONDERADO · EJECUTADO'): tot['ejec'] = {f: num(A.cell(r, c).value) for c, f in cols}
        if lab.startswith('CUMPLIMIENTO DE LA PROGRAMACIÓN'): tot['cumpl'] = {f: num(A.cell(r, c).value) for c, f in cols}
    return wb, fecha, cab, cols, acts, tot

def leer_personal(wb, fecha):
    P = wb['PERSONAL Y RESTRICCIONES']
    cols = {iso(P.cell(5, c).value): c for c in range(9, 27) if iso(P.cell(5, c).value)}
    fila = lambda txt: next((r for r in range(1, P.max_row + 1) if str(P.cell(r, 2).value or '').strip().startswith(txt)), None)
    serie = {}
    rA, rI, rT, rY, rS = fila('PRESENTES'), fila('… ingenieros'), fila('… técnicos'), fila('… ayudantes'), fila('PRESENTES (supervisión)')
    for f, c in cols.items():
        v = lambda r: (P.cell(r, c).value if r else None)
        serie[f] = dict(axon=v(rA), ing=v(rI), tec=v(rT), ayu=v(rY), sup=v(rS))
    # checklist del día (requisitos para iniciar la jornada)
    ini = fila('B · CHECKLIST'); chk = {}
    if ini:
        for r in range(ini + 3, ini + 12):
            req = P.cell(r, 2).value
            if not req or str(req).startswith('C ·'): break
            chk[str(req).strip()] = {f: P.cell(r, c).value for f, c in cols.items() if P.cell(r, c).value}
    hab, res = [], []
    for r in range(1, P.max_row + 1):
        a = P.cell(r, 2).value
        if isinstance(a, str) and a.startswith('H-') and P.cell(r, 3).value:
            hab.append(dict(id=a, nombre=sin_nombres(P.cell(r, 3).value), tipo=P.cell(r, 4).value, act=str(P.cell(r, 5).value or '').strip() or None,
                            resp=P.cell(r, 6).value, requerida=iso(P.cell(r, 9).value), estado=P.cell(r, 12).value,
                            cumplido=iso(P.cell(r, 16).value), alerta=P.cell(r, 19).value,
                            obs=sin_nombres(P.cell(r, 22).value), conenhua=sin_nombres(P.cell(r, 28).value)))
        if isinstance(a, str) and a.startswith('R-') and P.cell(r, 3).value:
            res.append(dict(id=a, desc=sin_nombres(P.cell(r, 3).value), cat=P.cell(r, 4).value, fecha=iso(P.cell(r, 5).value),
                            act=str(P.cell(r, 6).value or '').strip(), resp=P.cell(r, 9).value, hh=num(P.cell(r, 12).value, 1),
                            compromiso=iso(P.cell(r, 15).value), cierre=iso(P.cell(r, 18).value), estado=P.cell(r, 21).value,
                            necesita=sin_nombres(P.cell(r, 24).value), conenhua=sin_nombres(P.cell(r, 28).value)))
    return serie, chk, hab, res

# ---------------------------------------------------------------- cronogramas de referencia (PDF)
MESES = {m: i + 1 for i, m in enumerate(['jan', 'feb', 'mar', 'apr', 'may', 'jun', 'jul', 'aug', 'sep', 'oct', 'nov', 'dec'])}
MESES.update({'ene': 1, 'abr': 4, 'ago': 8, 'set': 9, 'dic': 12})

def pdf_txt(path):
    return subprocess.run(['pdftotext', '-layout', path, '-'], capture_output=True, text=True, check=True).stdout.splitlines()

def cron_v8(path, anio=2026):
    L, out = pdf_txt(path), []
    dl = re.compile(r'(\d{1,2})-([A-Za-z]{3})\s+(\d+)\s*$')
    for k, l in enumerate(L):
        m = re.match(r'^\s*(\d+\.\d+)\s+(.*?)\s*(\d+)%\s*$', l)
        if not m: continue
        item, nom, pct = m.group(1), m.group(2).strip(), int(m.group(3))
        nxt = L[k + 1] if k + 1 < len(L) else ''
        md = dl.search(nxt)
        if not md: continue
        ini = dt.date(anio, MESES[md.group(2).lower()], int(md.group(1))); dur = int(md.group(3))
        if not nom: nom = (L[k - 1].strip() + ' ' + nxt[:md.start()].strip()).strip()
        nom = re.sub(r'L-\s+1123', 'L-1123', nom)
        out.append(dict(item=item, nombre=nom, ini=ini.isoformat(), fin=(ini + dt.timedelta(days=dur - 1)).isoformat(), dias=dur,
                        fase='Fase 1 · inicio y documentación' if item.startswith('1.') else 'Fase 2 · ejecución', pct_doc=pct))
    return out

def cron_axon_r1(path):
    L, out = pdf_txt(path), []
    rx = re.compile(r'^\s*(\d{1,2})?\s*(.*?)\s{2,}(Inducción|Preparación|Montaje y tendido|Pruebas e integración|Cierre)\s+(\d{2})/(\d{2})/(\d{4})\s+(\d{2})/(\d{2})/(\d{4})\s+(\d+)\s*$')
    for k, l in enumerate(L):
        m = rx.match(l)
        if not m: continue
        n, desc, fase = m.group(1), m.group(2).strip(), m.group(3)
        ini = dt.date(int(m.group(6)), int(m.group(5)), int(m.group(4))); fin = dt.date(int(m.group(9)), int(m.group(8)), int(m.group(7)))
        if not desc: desc = (L[k - 1].strip() + ' ' + (L[k + 1].strip() if k + 1 < len(L) else '')).strip()
        out.append(dict(item=n, nombre=desc, fase=fase, ini=ini.isoformat(), fin=fin.isoformat(), dias=int(m.group(10))))
    return out

# ---------------------------------------------------------------- relación actividad ↔ modelo 3D
# «TABLERO / UBICACIÓN» del control de AXON → zona del modelo (ninguna ubicación sale de un plano: son SUPUESTAS y así se
# dicen en la ficha; ver AX_ZONAS en index.html)
ZONAS_DE = {'General': [], 'Almacén': ['almacen'], 'Patio': ['patio'], 'Patio → CPM': ['patio', 'cpm'], 'CPM L-1123': ['cpm'],
            'SSAA': ['ssaa'], 'SCADA SE': ['scada'], 'Tablero RF': ['t19', 'almacen']}

def serie_acum(d, fechas, hasta=None):
    """% acumulado por fecha: lo vacío antes del primer dato es 0; después, el último valor (la hoja deja en blanco los días
    posteriores al 100 %). Con `hasta`, None después de esa fecha (el real no se conoce)."""
    out, ult = [], 0.0
    for f in fechas:
        if hasta and f > hasta: out.append(None); continue
        v = d.get(f)
        if v is not None: ult = v
        out.append(round(ult, 6))
    return out

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--reportes', required=True); ap.add_argument('--v8'); ap.add_argument('--axr1'); ap.add_argument('--fotos')
    ap.add_argument('--hasta', help='fecha final de la ventana 4D (por defecto, el último día en que cambia la línea base)')
    ap.add_argument('--out', required=True)
    a = ap.parse_args()
    archivos = sorted(glob.glob(os.path.join(a.reportes, 'axon_*.xlsx')))
    if not archivos: raise SystemExit('sin controles de AXON en ' + a.reportes)
    por_dia = {}
    for p in archivos:
        wb, fecha, cab, cols, acts, tot = leer_avance(p)
        serie, chk, hab, res = leer_personal(wb, fecha)
        por_dia[fecha] = dict(archivo=os.path.basename(p), cab=cab, cols=cols, acts=acts, tot=tot, personal=serie, chk=chk, hab=hab, res=res)
    corte = max(por_dia); U = por_dia[corte]
    fechas_todas = [f for _, f in U['cols']]
    tp = U['tot']['prog']
    fin_lb = max(f for f in fechas_todas if tp.get(f) is not None and tp[f] != tp.get(fechas_todas[fechas_todas.index(f) - 1]) ) if tp else corte
    hasta = a.hasta or max(fin_lb, corte)
    fechas = [f for f in fechas_todas if f <= hasta]
    acts, ajustes = [], []
    for x in U['acts']:
        p = serie_acum(x['_prog'], fechas); e = serie_acum(x['_ejec'], fechas, hasta=corte)
        # Un % acumulado no retrocede (decisión del usuario, 25-Oct-26: «la curva S no puede bajar; puede estancarse»).
        # Si el control trae un valor menor que el máximo anterior (p. ej. 0 el día que no se trabajó), se toma como
        # «sin avance ese día» y se guarda el dato original en `ajustes` para mostrarlo.
        mx = 0.0
        for i, v in enumerate(e):
            if v is None: continue
            if v < mx - 1e-9:
                ajustes.append(dict(f=fechas[i], n=x['n'], nombre=x['nombre'], peso=x['peso'], axon=v, usado=mx))
                e[i] = mx
            mx = max(mx, e[i])
        con = [f for f in fechas if (x['_prog'].get(f) or 0) > 0]
        ini_lb = x['ini'] or (con[0] if con else None)
        fin_lb_a = x['fin'] or (con[-1] if con else None)
        # causa / justificación declaradas en el control de cada día
        hist = {}
        for f, D in por_dia.items():
            y = next((z for z in D['acts'] if z['n'] == x['n']), None)
            if y and (y['causa'] or y['justif']): hist[f] = dict(causa=y['causa'], justif=y['justif'], estado=y['estado'])
        acts.append(dict(n=x['n'], nombre=x['nombre'], ubic=x['ubic'], zonas=ZONAS_DE.get(x['ubic'], []), peso=x['peso'], ritmo=x['ritmo'],
                         seccion=x['seccion'], grupo=x['grupo'], alcance=x['alcance'], ini=ini_lb, fin=fin_lb_a,
                         prog=p, ejec=e, estado=x['estado'], pv=x['pv'],
                         ev=(e[fechas.index(corte)] if corte in fechas and any(j['n'] == x['n'] for j in ajustes) else x['ev']),
                         ev_axon=(x['ev'] if any(j['n'] == x['n'] for j in ajustes) else None), causas=hist))
    tot = dict(prog=serie_acum(U['tot']['prog'], fechas), ejec=serie_acum(U['tot'].get('ejec', {}), fechas, hasta=corte),
               cumpl=[U['tot'].get('cumpl', {}).get(f) for f in fechas])
    dias_aj = sorted({j['f'] for j in ajustes})
    if dias_aj:
        tot['ejec_axon'] = list(tot['ejec'])
        tot['cumpl_axon'] = list(tot['cumpl'])
        for f in dias_aj:
            i = fechas.index(f)
            tot['ejec'][i] = round(tot['ejec'][i] + sum((j['peso'] or 0) * (j['usado'] - j['axon']) for j in ajustes if j['f'] == f), 6)
            # cumplimiento del día = promedio de (avance ejecutado ÷ avance programado) de las actividades programadas ese día
            rs = []
            for x in acts:
                dp = (x['prog'][i] or 0) - ((x['prog'][i - 1] or 0) if i else 0)
                if dp > 1e-9:
                    de = (x['ejec'][i] or 0) - ((x['ejec'][i - 1] or 0) if i else 0)
                    rs.append(min(1.0, max(0.0, de / dp)))
            tot['cumpl'][i] = round(sum(rs) / len(rs), 6) if rs else None
        # un día en blanco después del error arrastra el valor bajo: queda como ajuste de ese día y se corrige igual
    # reporte de cada día: el control de ese día (o, si no lo hay, lo que la serie del último control dice de ese día)
    dias = {}
    for i, f in enumerate(fechas):
        if f > corte: break
        D = por_dia.get(f)
        pers = U['personal'].get(f) or {}
        hechas = [x['n'] for x in acts if x['ejec'][i] is not None and x['ejec'][i] > (x['ejec'][i - 1] if i else 0) + 1e-6]
        dia = dict(personal=pers, hechas=hechas,
                   checklist={k: v.get(f) for k, v in U['chk'].items() if v.get(f)},
                   restr_abiertas=[r['id'] for r in U['res'] if r['fecha'] and r['fecha'] <= f and not (r['cierre'] and r['cierre'] < f)],
                   restr_nuevas=[r['id'] for r in U['res'] if r['fecha'] == f])
        if D:
            c = D['cab']
            dia.update(reporte=c.get('Reporte N°'), dia_plazo=c.get('Día del plazo'), prog=num(c.get('Programado al corte')),
                       ejec=num(c.get('Ejecutado al corte')), spi=num(c.get('SPI (ejec. / prog.)'), 3), ejec_dia=num(c.get('Ejecutado del día (pts)'), 2),
                       atrasadas=c.get('Actividades atrasadas'), restr_n=c.get('Restricciones abiertas'), hab_vencidos_n=c.get('Habilitantes vencidos'),
                       dias_fin=c.get('Días al fin del plan'), archivo=D['archivo'],
                       hab_vencidos=[h['id'] for h in D['hab'] if h['alerta'] == 'Vencido'])
        if f in dias_aj and dia.get('ejec') is not None:
            dia.update(ejec_axon=dia['ejec'], spi_axon=dia.get('spi'), ejec_dia_axon=dia.get('ejec_dia'),
                       ejec=tot['ejec'][i], spi=round(tot['ejec'][i] / tot['prog'][i], 3) if tot['prog'][i] else None,
                       ejec_dia=round((tot['ejec'][i] - (tot['ejec'][i - 1] if i else 0)) * 100, 2),
                       ajustes=[dict(n=j['n'], nombre=j['nombre'], axon=j['axon'], usado=j['usado']) for j in ajustes if j['f'] == f])
        dias[f] = dia
    # Actividades que AXON agrega en un control posterior (p. ej. la 10.2 del 26-Oct, fechada el 19-Oct): cambian los pesos y
    # recalculan días ya publicados. El tablero muestra la serie del control vigente y dice qué decía el control anterior.
    rebase, fcs = [], sorted(por_dia)
    base = lambda D: round(sum((z.get('dias_peso') or 0) for z in D['acts']), 3)
    for ic in range(1, len(fcs)):
        fc, fp = fcs[ic], fcs[ic - 1]
        ns_prev = {z['n'] for z in por_dia[fp]['acts']}
        nuevas = [z for z in por_dia[fc]['acts'] if z['n'] not in ns_prev]
        if not nuevas: continue
        Tp = por_dia[fp]['tot']; cambios = []
        for i, f in enumerate(fechas):
            if f > fp: break
            pa, ea = Tp.get('prog', {}).get(f), Tp.get('ejec', {}).get(f)
            pn, en = tot['prog'][i], tot['ejec'][i]
            if (pa is not None and pn is not None and abs(pa - pn) > 0.0005) or (ea is not None and en is not None and abs(ea - en) > 0.0005):
                cambios.append(dict(f=f, prog_antes=num(pa), ejec_antes=num(ea), prog=pn, ejec=en))
        rebase.append(dict(control=fc, previo=fp, base_antes=base(por_dia[fp]), base=base(por_dia[fc]), cambios=cambios,
                           nuevas=[dict(n=z['n'], nombre=z['nombre'], peso=z['peso'], dias_peso=z['dias_peso'], ini=z['ini'], fin=z['fin'])
                                   for z in nuevas]))
    data = dict(
        titulo='Habilitación del nuevo tablero CPM L-1123 (=E01+WR01) · SE Uchucchacua',
        contratista='AXON Group (AX-26138 · OS BVN 5800003219)', supervision='CONENHUA', propietario='Buenaventura',
        linea_base='Control diario de AXON (sigue el cronograma V8 de CONENHUA)',
        fuente=f'Control diario de AXON, reportes N.º {min(d["cab"].get("Reporte N°") for d in por_dia.values())}–{U["cab"].get("Reporte N°")} '
               f'({", ".join(sorted(f[8:10] for f in por_dia))}-Oct-26) · hojas AVANCE y PERSONAL Y RESTRICCIONES',
        corte=corte, fechas=fechas, total=tot, acts=acts, dias=dias, ajustes=ajustes, rebase=rebase,
        habilitantes=U['hab'], restricciones=U['res'],
        cronogramas=dict(v8=cron_v8(a.v8) if a.v8 else [], axon_r1=cron_axon_r1(a.axr1) if a.axr1 else []),
        fotos=relacionar_fotos(json.load(open(a.fotos, encoding='utf-8')).get('fotos', []) if a.fotos and os.path.exists(a.fotos) else [], acts, dias))
    open(a.out, 'w', encoding='utf-8').write('// BIM 4D de la SE Uchucchacua · generado con gen/uchu_build.py desde los controles diarios de AXON y los '
                                            'cronogramas V8 (CONENHUA) y Rev. 1 (AXON); no editar a mano.\nexport const U4D = '
                                            + json.dumps(data, ensure_ascii=False, separators=(',', ':')) + ';\n')
    print(f'OK {a.out} · corte {corte} · {len(fechas)} días ({fechas[0]} → {fechas[-1]}) · {len(acts)} actividades · '
          f'{len(data["fotos"])} fotos · V8 {len(data["cronogramas"]["v8"])} · AXON R1 {len(data["cronogramas"]["axon_r1"])}')
    i = fechas.index(corte)
    print(f'  al corte: programado {tot["prog"][i]} · ejecutado {tot["ejec"][i]} (cabecera: {U["cab"].get("Programado al corte")} / {U["cab"].get("Ejecutado al corte")})')

if __name__ == '__main__':
    main()
